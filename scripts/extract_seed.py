"""Extract 2018 planting documents into seed JSON. Usage: python3 -I extract_seed.py <schedule.xlsx> <treelist.xlsx> <out.json>"""
import json, re, sys
import openpyxl

sched, trees_path, out = sys.argv[1:4]
BED_IDS = {
    'The Dingley Dell (including the Mimosa and Catalpa Beds)': 'bed-dingley-dell',
    'Roadside Bed East Side': 'bed-roadside-east',
    'Roadside Bed West Side': 'bed-roadside-west',
    'Border Behind Reception': 'bed-border-reception',
}
def slug(s): return re.sub(r'[^a-z0-9]+', '-', s.lower()).strip('-')

wb = openpyxl.load_workbook(sched, data_only=True)
ws = wb['Shrubs and herbaceous']
drifts, taxa, bed = [], {}, None
for row in ws.iter_rows(values_only=True):
    cells = [c for c in row if c is not None]
    if not cells or not isinstance(cells[0], str): continue
    t = cells[0].strip()
    if t in BED_IDS: bed = BED_IDS[t]; continue
    m = re.match(r'^(\d+)\s*No\.?\s*(.+)$', t)
    if not (m and bed): continue
    qty, rest = int(m.group(1)), m.group(2).strip()
    note = None
    nm = re.search(r'\((.+)\)\s*$', rest)
    if nm: note, rest = nm.group(1), rest[:nm.start()].strip()
    name = re.sub(r'\s+', ' ', rest).strip().replace('Sarcococcoca', 'Sarcococca').replace('Perovksia', 'Perovskia').replace('Dodonea', 'Dodonaea').replace('Fucshia', 'Fuchsia').replace('Eleagnus', 'Elaeagnus').replace('Honerine', 'Honorine').replace('suphureum', 'sulphureum')
    tid = 'taxon-' + slug(name)
    taxa.setdefault(tid, {'id': tid, 'name': name, 'genus': name.split()[0]})
    pot = cells[1] if len(cells) > 1 else None
    drifts.append({'id': f'drift-{bed[4:]}-{slug(name)}', 'bedId': bed, 'taxonId': tid, 'quantity': qty,
                   'potSize': str(pot).split('@')[0].strip() if pot else None, 'note': note})

trees = []
tw = openpyxl.load_workbook(trees_path, data_only=True).active
for row in tw.iter_rows(min_row=4, values_only=True):
    if not row[0] or not row[5]: continue
    name = row[0].strip(); tid = 'taxon-' + slug(name)
    taxa.setdefault(tid, {'id': tid, 'name': name, 'genus': name.split()[0], 'common': row[1].strip(' "')})
    for i in range(int(row[5])):
        trees.append({'id': f'tree-{slug(name)}-{slug(row[2])}-{i+1}', 'taxonId': tid, 'location': row[2], 'size': row[3]})

json.dump({'drifts': drifts, 'taxa': sorted(taxa.values(), key=lambda t: t['name']), 'trees2018': trees},
          open(out, 'w'), indent=1, ensure_ascii=False)
print(len(drifts), 'drifts', len(taxa), 'taxa', len(trees), 'trees', sum(d['quantity'] for d in drifts), 'plants')
